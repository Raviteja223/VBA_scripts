from __future__ import annotations

from pathlib import Path
from typing import Any
import unicodedata

import pandas as pd
from openpyxl import load_workbook


def normalize_header(value: Any) -> str:
    """
    Normalize Excel column headers without changing their business meaning.

    Handles:
    - leading/trailing spaces
    - repeated spaces
    - non-breaking spaces
    - zero-width characters
    - BOM characters
    - Unicode compatibility differences

    Examples:
        " Statut "         -> "Statut"
        "Date\xa0de retour" -> "Date de retour"
        "Statut\u200b"      -> "Statut"
    """

    if value is None:
        return ""

    text = str(value)

    # Normalize Unicode representation.
    text = unicodedata.normalize("NFKC", text)

    # Common invisible characters found in copied/imported Excel headers.
    text = (
        text
        .replace("\xa0", " ")      # non-breaking space
        .replace("\u2007", " ")    # figure space
        .replace("\u202f", " ")    # narrow non-breaking space
        .replace("\u200b", "")     # zero-width space
        .replace("\u200c", "")     # zero-width non-joiner
        .replace("\u200d", "")     # zero-width joiner
        .replace("\ufeff", "")     # BOM / zero-width no-break space
    )

    # Collapse multiple whitespace characters into one normal space.
    text = " ".join(text.split())

    return text.strip()


def _normalize_dataframe_columns(df: pd.DataFrame) -> pd.DataFrame:
    """
    Normalize all dataframe column names and ensure normalization
    does not create duplicate headers.
    """

    original_columns = list(df.columns)

    normalized_columns = [
        normalize_header(column)
        for column in original_columns
    ]

    # Detect duplicate names created by normalization.
    duplicates = sorted(
        {
            column
            for column in normalized_columns
            if normalized_columns.count(column) > 1
        }
    )

    if duplicates:
        details = []

        for duplicate in duplicates:
            originals = [
                repr(original)
                for original, normalized in zip(
                    original_columns,
                    normalized_columns,
                )
                if normalized == duplicate
            ]

            details.append(
                f"{duplicate!r} <- {', '.join(originals)}"
            )

        raise ValueError(
            "Duplicate Excel column names were created after "
            "header normalization:\n"
            + "\n".join(details)
        )

    result = df.copy()
    result.columns = normalized_columns

    return result


def sheet_names(path: Path | str) -> list[str]:
    """
    Return all worksheet names from an Excel workbook.

    The workbook is opened read-only and is not modified.
    """

    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(
            f"Excel workbook not found: {path}"
        )

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=False,
    )

    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def read_sheet(
    path: Path | str,
    sheet: str,
    **kwargs: Any,
) -> pd.DataFrame:
    """
    Read one worksheet into a pandas DataFrame and normalize
    its column headers.

    Extra arguments such as nrows, header, dtype, usecols, etc.
    are passed directly to pandas.read_excel().
    """

    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(
            f"Excel workbook not found: {path}"
        )

    available_sheets = sheet_names(path)

    if sheet not in available_sheets:
        raise KeyError(
            f"Worksheet {sheet!r} was not found in {path.name}. "
            f"Available sheets: {available_sheets}"
        )

    df = pd.read_excel(
        path,
        sheet_name=sheet,
        engine="openpyxl",
        **kwargs,
    )

    return _normalize_dataframe_columns(df)


def read_configured(
    path: Path | str,
    sheets: list[str],
) -> dict[str, pd.DataFrame]:
    """
    Read only the configured worksheets.

    Returns:
        {
            "Remarketing": <DataFrame>,
            "Mapping": <DataFrame>,
            ...
        }
    """

    path = Path(path)

    if not sheets:
        raise ValueError(
            "No worksheet names were supplied to read_configured()."
        )

    available = set(sheet_names(path))

    missing = [
        sheet
        for sheet in sheets
        if sheet not in available
    ]

    if missing:
        raise KeyError(
            "Configured worksheet(s) do not exist in the workbook: "
            + ", ".join(missing)
        )

    result: dict[str, pd.DataFrame] = {}

    for sheet in sheets:
        result[sheet] = read_sheet(
            path=path,
            sheet=sheet,
        )

    return result
